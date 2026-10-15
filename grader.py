import openpyxl
from openpyxl.styles import Font, PatternFill

# 1. Create workbook and set sheet title
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Semester_Results"

# 2. Header row
headers = ["Roll No", "Student Name", "Lab Marks (out of 50)",
           "Theory Marks (out of 50)", "Total Marks (out of 100)", "Result"]
ws.append(headers)

# 3. Sample student data: Roll No, Name, Lab Marks, Theory Marks
students = [
    (1, "Aarav Shah", 42, 38),
    (2, "Priya Mehta", 35, 30),
    (3, "Rohan Desai", 20, 15),
    (4, "Isha Patel", 48, 45),
    (5, "Kabir Nair", 25, 10),
]

for roll, name, lab, theory in students:
    total = lab + theory
    result = "Pass" if total >= 40 else "Fail"
    ws.append([roll, name, lab, theory, total, result])

# 4. Styling the header row
header_font = Font(bold=True)
header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill

# 5. Compute class average and add summary row
total_col = 5  # "Total Marks" is column E (1-indexed: 5)
num_students = len(students)
avg_total = sum(row[2] + row[3] for row in students) / num_students

summary_row = ["", "", "", "Class Average:", round(avg_total, 2), ""]
ws.append(summary_row)

# Bold the summary row too
for cell in ws[ws.max_row]:
    cell.font = Font(bold=True)

# Save workbook
wb.save("AIML_Batch_Grades.xlsx")
print("Workbook 'AIML_Batch_Grades.xlsx' created successfully.")